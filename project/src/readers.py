import csv
from typing import Any

from openpyxl import load_workbook


def _build_transaction(row: dict[str, Any]) -> dict[str, Any]:
    """Превращает плоскую строку таблицы в формат, как в operations.json."""
    return {
        "id": row.get("id"),
        "state": row.get("state"),
        "date": row.get("date"),
        "operationAmount": {
            "amount": str(row.get("amount")),
            "currency": {
                "name": row.get("currency_name"),
                "code": row.get("currency_code"),
            },
        },
        "description": row.get("description"),
        "from": row.get("from") or None,
        "to": row.get("to") or None,
    }


def read_csv(path: str) -> list[dict[str, Any]]:
    """Читает транзакции из CSV-файла (разделитель ";"). При проблеме возвращает []."""
    try:
        with open(path, encoding="utf-8", newline="") as file:
            return [_build_transaction(dict(row)) for row in csv.DictReader(file, delimiter=";")]
    except (OSError, csv.Error):
        return []


def read_xlsx(path: str) -> list[dict[str, Any]]:
    """Читает транзакции из XLSX-файла (первая строка — заголовки). При проблеме возвращает []."""
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (OSError, ValueError, KeyError):
        return []

    sheet = workbook.active
    if sheet is None:
        return []

    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        return []

    names = [str(name) for name in header]
    return [_build_transaction(dict(zip(names, row))) for row in rows if any(row)]
